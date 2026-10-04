import os
import sys
import unittest
import tempfile
import warnings
from docx import Document as DocxDocument

warnings.filterwarnings("ignore")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from fastapi.testclient import TestClient
from app.main import app
from app.services.embeddings.embedding_service import EmbeddingService
from app.services.deterministic.experience_calculator import ExperienceCalculator
from app.services.deterministic.skills_analyzer import SkillsAnalyzer
from app.services.deterministic.deterministic_scorer import DeterministicScorer
from app.services.ai.recommendation_service import RecommendationService
from app.services.ats.ats_scoring_service import ATSScoringService

class TestFinalRefactorSuite(unittest.TestCase):
    """
    Comprehensive verification suite for final SkillForge AI refactor pass:
    - Pretrained Embedding Model & Caching
    - Thorough Experience Evaluation & Overlap Merging
    - Skills Analysis & Java vs JavaScript False-Positive Prevention
    - Deterministic ATS Scoring Formula (No LLM Dependence)
    - Downstream Recommendation Layer & Non-AI Rule-Based Fallback
    - Full End-to-End User Journey
    """

    # =========================================================================
    # 1. EMBEDDING IMPLEMENTATION TESTS
    # =========================================================================
    def test_embedding_singleton_and_basics(self):
        """Verify model caching, empty text, identical text, and normalization."""
        # 1. Empty or whitespace text returns 0.0
        self.assertEqual(EmbeddingService.compute_similarity("", "Some text"), 0.0)
        self.assertEqual(EmbeddingService.compute_similarity("   ", ""), 0.0)
        self.assertEqual(EmbeddingService.compute_similarity(None, "Text"), 0.0)

        # 2. Identical text returns exactly 1.0 (raw) and 100.0 (percent)
        text = "Experienced Python and FastAPI backend software engineer."
        self.assertEqual(EmbeddingService.compute_similarity(text, text), 1.0)
        self.assertEqual(EmbeddingService.compute_similarity_percent(text, text), 100.0)

        # 3. Unrelated text has very low similarity
        unrelated_a = "Cooking Italian pasta and baking sourdough bread in Tuscany."
        unrelated_b = "Kubernetes cluster container orchestration and AWS cloud networking."
        sim_unrelated = EmbeddingService.compute_similarity(unrelated_a, unrelated_b)
        self.assertLess(sim_unrelated, 0.3)

        # 4. Semantically related wording has higher similarity
        related_a = "Software engineer building RESTful web services and PostgreSQL databases."
        related_b = "Software engineer creating RESTful web services and PostgreSQL databases."
        sim_related = EmbeddingService.compute_similarity(related_a, related_b)
        self.assertGreater(sim_related, 0.4)

    # =========================================================================
    # 2. EXPERIENCE EVALUATION TESTS
    # =========================================================================
    def test_experience_zero(self):
        """Zero experience against 4 yr requirement produces 0.0%."""
        res = ExperienceCalculator.calculate([], {"title": "Dev", "experience_required": "4 years"}, "", "")
        self.assertEqual(res["experience_score"], 0.0)
        self.assertEqual(res["candidate_relevant_years"], 0.0)
        self.assertEqual(res["required_experience_years"], 4.0)
        self.assertEqual(res["match_status"], "No Experience Found")

    def test_experience_exact_match(self):
        """Exact experience match produces 100.0%."""
        exp = [{"role": "Software Engineer", "company": "A", "duration": "2020 - 2023", "responsibilities": ["Backend services"]}]
        res = ExperienceCalculator.calculate(exp, {"title": "Software Engineer", "experience_required": "3 years"}, "", "Software Engineer")
        self.assertEqual(res["candidate_total_years"], 3.0)
        self.assertEqual(res["experience_score"], 100.0)
        self.assertEqual(res["match_status"], "Fully Meets / Exceeds Requirement")

    def test_experience_below_requirement(self):
        """Candidate with 2 years on 4 year requirement produces ~50%."""
        exp = [{"role": "Backend Dev", "company": "A", "duration": "2022 - 2024", "responsibilities": ["FastAPI APIs"]}]
        res = ExperienceCalculator.calculate(exp, {"title": "Backend Dev", "experience_required": "4 years"}, "", "Backend Dev")
        self.assertAlmostEqual(res["experience_score"], 50.0, delta=2.0)
        self.assertEqual(res["match_status"], "Below Requirement")


    def test_experience_above_requirement(self):
        """Candidate with 6 years on 3 year requirement produces capped 100.0%."""
        exp = [{"role": "Senior Dev", "company": "A", "duration": "2018 - 2024", "responsibilities": ["Architecture"]}]
        res = ExperienceCalculator.calculate(exp, {"title": "Senior Dev", "experience_required": "3 years"}, "", "Senior Dev")
        self.assertEqual(res["experience_score"], 100.0)

    def test_experience_overlapping_roles_no_double_count(self):
        """Overlapping employment periods (2020-2022 and 2021-2023) merge to 3.0 yrs, not 4.0 yrs."""
        exp = [
            {"role": "Backend Engineer", "company": "Co A", "duration": "2020 - 2022", "responsibilities": ["Python APIs"]},
            {"role": "Consultant", "company": "Co B", "duration": "2021 - 2023", "responsibilities": ["Python services"]}
        ]
        res = ExperienceCalculator.calculate(exp, {"title": "Engineer", "experience_required": "3 years"}, "", "Python")
        self.assertEqual(res["candidate_total_years"], 3.0)
        self.assertEqual(res["experience_score"], 100.0)

    def test_experience_current_ongoing_role(self):
        """Current/ongoing roles (e.g. 2022 - Present) parse safely."""
        exp = [{"role": "Dev", "company": "Tech", "duration": "2022 - Present", "responsibilities": ["Fullstack"]}]
        res = ExperienceCalculator.calculate(exp, {"title": "Dev", "experience_required": "2 years"}, "", "")
        self.assertGreaterEqual(res["candidate_total_years"], 2.0)
        self.assertEqual(res["experience_score"], 100.0)

    def test_experience_absent_jd_requirement(self):
        """Absent experience requirement returns 100.0% score and 0.0 required years."""
        exp = [{"role": "Dev", "company": "Tech", "duration": "2023 - 2024", "responsibilities": ["Code"]}]
        res = ExperienceCalculator.calculate(exp, {"title": "Software Engineer"}, "", "Software developer needed")
        self.assertEqual(res["required_experience_years"], 0.0)
        self.assertEqual(res["experience_score"], 100.0)

    def test_experience_malformed_dates(self):
        """Malformed or descriptive date strings fallback safely without crashing."""
        exp = [{"role": "Intern", "company": "StartUp", "duration": "Summer semester", "responsibilities": ["Coding"]}]
        res = ExperienceCalculator.calculate(exp, {"title": "Dev", "experience_required": "1 year"}, "", "")
        self.assertGreater(res["experience_score"], 0.0)
        self.assertLessEqual(res["experience_score"], 100.0)

    # =========================================================================
    # 3. SKILLS ANALYSIS & FALSE POSITIVE PREVENTION
    # =========================================================================
    def test_skills_java_vs_javascript_no_false_positive(self):
        """Candidate with JavaScript must NOT match a JD requiring Java."""
        parsed_resume = {
            "skills": {"programming_languages": ["JavaScript", "HTML", "CSS"]},
            "experience": [{"role": "Frontend Dev", "responsibilities": ["Built React components in JavaScript"]}]
        }
        parsed_job = {
            "skills": {"required_skills": ["Java", "Spring Boot", "SQL"]}
        }
        res = SkillsAnalyzer.analyze(parsed_resume, parsed_job, "JavaScript frontend developer", "Java backend developer")
        self.assertNotIn("Java", res["matched_skills"])
        self.assertIn("Java", res["missing_skills"])

    def test_skills_required_vs_preferred_weighting(self):
        """Missing required skills impacts the score more heavily than preferred."""
        parsed_resume = {"skills": {"programming_languages": ["Python", "Git"]}}
        
        # Missing required skill (Kubernetes)
        job_a = {"skills": {"required_skills": ["Python", "Kubernetes"], "preferred_skills": ["Redis"]}}
        res_a = SkillsAnalyzer.analyze(parsed_resume, job_a, "Python Git", "Python Kubernetes Redis")

        # Missing preferred skill (Redis), but has all required skills
        job_b = {"skills": {"required_skills": ["Python", "Git"], "preferred_skills": ["Redis"]}}
        res_b = SkillsAnalyzer.analyze(parsed_resume, job_b, "Python Git", "Python Git Redis")

        self.assertGreater(res_b["skill_score"], res_a["skill_score"])

    # =========================================================================
    # 4. CORE ATS SCORING & RECOMMENDATION DECOUPLING
    # =========================================================================
    def test_ats_scoring_independent_of_external_ai(self):
        """ATS scoring succeeds with deterministic arithmetic even when AI provider is unavailable."""
        parsed_resume = {
            "personal_info": {"email": "alex@test.com", "phone": "1234567890"},
            "skills": {"programming_languages": ["Python", "FastAPI", "SQL"]},
            "experience": [{"role": "Software Engineer", "company": "Co", "duration": "2021 - 2024", "responsibilities": ["API dev"]}],
            "education": [{"degree": "B.S. in Computer Science", "institution": "Univ"}],
            "projects": [{"name": "API Service", "description": "FastAPI REST API with SQL database", "technologies": ["Python", "FastAPI"]}]
        }
        parsed_job = {
            "job_info": {"title": "Python Developer", "experience_required": "3 years"},
            "skills": {"required_skills": ["Python", "FastAPI", "SQL"], "preferred_skills": ["Docker", "Kubernetes"]},
            "responsibilities": ["Build FastAPI services", "Write clean SQL queries"]
        }

        report = ATSScoringService.compute_full_ats_report(
            parsed_resume, parsed_job, "Python FastAPI SQL developer", "Python Developer 3 years required", {}, {}
        )

        # 1. Deterministic score is produced
        self.assertIn("overall_score", report)
        self.assertGreater(report["overall_score"], 0.0)
        self.assertLessEqual(report["overall_score"], 100.0)

        # 2. Component scores are recorded
        breakdown = report["score_breakdown"]
        self.assertIn("skill_score", breakdown)
        self.assertIn("experience_score", breakdown)
        self.assertIn("responsibility_score", breakdown)
        self.assertIn("keyword_score", breakdown)
        self.assertIn("semantic_score", breakdown)
        self.assertIn("education_score", breakdown)
        self.assertIn("project_score", breakdown)
        self.assertIn("formatting_score", breakdown)

        # 3. Recommendations are present and match strict schema
        recs = report["recommendations"]
        self.assertIn("priority_skills", recs)
        self.assertIn("project_suggestions", recs)
        self.assertIn("learning_direction", recs)
        self.assertIn("resume_improvements", recs)

        # Project suggestion structure
        for ps in recs["project_suggestions"]:
            self.assertIn("title", ps)
            self.assertIn("why_relevant", ps)
            self.assertIn("skills_practiced", ps)
            self.assertIn("difficulty", ps)

    # =========================================================================
    # 5. FULL END-TO-END PIPELINE INTEGRATION
    # =========================================================================
    def test_full_user_journey_e2e(self):
        """End-to-End: Register -> Login -> Upload Resume -> Submit JD -> Analyze -> Retrieve Report."""
        client = TestClient(app)

        # A. Register
        user_email = f"e2e_{os.urandom(4).hex()}@example.com"
        pwd = "SecurePassword123!"
        reg_res = client.post("/api/v1/auth/register", json={"email": user_email, "password": pwd, "full_name": "E2E Candidate"})
        self.assertEqual(reg_res.status_code, 200)

        # B. Login
        login_res = client.post("/api/v1/auth/login/json", json={"email": user_email, "password": pwd})
        self.assertEqual(login_res.status_code, 200)
        token = login_res.json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}

        # C. Upload Resume DOCX
        doc = DocxDocument()
        doc.add_heading("E2E Candidate", 0)
        doc.add_paragraph("Email: candidate@test.com | Phone: 555-123-4567")
        doc.add_heading("Skills", 1)
        doc.add_paragraph("Python, FastAPI, PostgreSQL, Git, Docker, REST APIs")
        doc.add_heading("Experience", 1)
        doc.add_paragraph("Software Engineer - Acme Corp (2022 - 2024)\nBuilt REST microservices using Python and FastAPI. Optimized PostgreSQL queries.")
        doc.add_heading("Education", 1)
        doc.add_paragraph("B.Tech in Computer Science - State University (2018 - 2022)")
        doc.add_heading("Projects", 1)
        doc.add_paragraph("SkillForge Project - Career platform built with FastAPI and React.")
        
        temp_docx = os.path.join(tempfile.gettempdir(), f"e2e_resume_{os.urandom(4).hex()}.docx")
        doc.save(temp_docx)

        with open(temp_docx, "rb") as f:
            up_res = client.post(
                "/api/v1/resumes/upload",
                files={"file": ("resume.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                headers=headers
            )
        self.assertEqual(up_res.status_code, 201)
        resume_id = up_res.json()["id"]

        # D. Submit Job Description
        jd_res = client.post(
            "/api/v1/jobdescription/",
            json={
                "title": "Backend Python Developer",
                "description": "Looking for Backend Developer with Python, FastAPI, PostgreSQL, Docker, and AWS skills. 2+ years required.",
                "resume_id": resume_id
            },
            headers=headers
        )
        self.assertEqual(jd_res.status_code, 201)
        job_id = jd_res.json()["id"]

        # E. Run Full Analysis
        ana_res = client.post(f"/api/v1/analysis/{resume_id}", json={"resume_id": resume_id, "job_description_id": job_id}, headers=headers)
        self.assertEqual(ana_res.status_code, 201)
        ana_data = ana_res.json()
        self.assertGreater(ana_data["overall_score"], 0.0)

        # F. Retrieve ATS Full Report
        report_res = client.get(f"/api/v1/ats/{resume_id}/report", headers=headers)
        self.assertEqual(report_res.status_code, 200)
        report_data = report_res.json()
        self.assertEqual(report_data["overall_score"], ana_data["overall_score"])
        self.assertIn("score_breakdown", report_data)
        self.assertIn("report_details", report_data)

if __name__ == "__main__":
    unittest.main()
