import os
import sys
import warnings
import tempfile

# Suppress starlette/httpx deprecation warning
warnings.filterwarnings("ignore", category=DeprecationWarning)

from fastapi.testclient import TestClient
# pyrefly: ignore [missing-import]
from docx import Document as DocxDocument

# Add current directory to path
sys.path.insert(0, os.path.dirname(__file__))

from app.main import app
from app.db.session import engine
from app.db.base import Base

# Ensure tables are created
Base.metadata.create_all(bind=engine)

client = TestClient(app)

def create_sample_resume_docx() -> str:
    """Create a sample DOCX resume file for testing and return its path."""
    doc = DocxDocument()
    doc.add_heading("John Doe", 0)
    doc.add_paragraph("Email: johndoe@example.com | Phone: +91-9876543210")
    doc.add_paragraph("GitHub: github.com/johndoe | LinkedIn: linkedin.com/in/johndoe")
    doc.add_heading("Summary", level=1)
    doc.add_paragraph("A passionate software developer skilled in Python, FastAPI, Angular, and SQL.")
    doc.add_heading("Skills", level=1)
    doc.add_paragraph("Python, JavaScript, TypeScript, FastAPI, Angular, PostgreSQL, Git, Docker, SQL, HTML, CSS")
    doc.add_heading("Work Experience", level=1)
    doc.add_paragraph("Software Engineer Intern - ABC Corp (2025-2026)\nDeveloped REST APIs using FastAPI and integrated PostgreSQL.")
    doc.add_heading("Education", level=1)
    doc.add_paragraph("B.Tech in Computer Science - XYZ University (2022-2026)")
    doc.add_heading("Projects", level=1)
    doc.add_paragraph("SkillForge AI - AI career mentorship platform built with Angular and FastAPI.")
    doc.add_heading("Certifications", level=1)
    doc.add_paragraph("AWS Cloud Practitioner, Python Data Science Certification")

    path = os.path.join(tempfile.gettempdir(), "sample_resume_test.docx")
    doc.save(path)
    return path

def test_full_flow():
    print("1. Testing Health Endpoint...")
    resp = client.get("/")
    assert resp.status_code == 200
    print("Health response:", resp.json())

    print("\n2. Testing User Registration...")
    reg_data = {
        "email": "testuser_weeks1to5@example.com",
        "password": "SecretPassword123!",
        "full_name": "Test Candidate"
    }
    resp = client.post("/api/v1/auth/register", json=reg_data)
    if resp.status_code == 400 and "already exists" in resp.json().get("detail", ""):
        print("User already registered, proceeding to login.")
    else:
        assert resp.status_code == 200, f"Registration failed: {resp.text}"
        print("Registration success!")

    print("\n3. Testing User Login (JSON)...")
    login_data = {
        "email": "testuser_weeks1to5@example.com",
        "password": "SecretPassword123!"
    }
    resp = client.post("/api/v1/auth/login/json", json=login_data)
    assert resp.status_code == 200, f"Login failed: {resp.text}"
    token = resp.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    print("Login successful! Got access token.")

    print("\n4. Testing GET /auth/me...")
    resp = client.get("/api/v1/auth/me", headers=headers)
    assert resp.status_code == 200
    print("Current User Me:", resp.json())

    print("\n5. Testing Resume DOCX Upload & Parsing...")
    sample_docx_path = create_sample_resume_docx()
    print(f"  Sample resume created at: {sample_docx_path}")

    with open(sample_docx_path, "rb") as f:
        files = {"file": ("john_doe_resume.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        resp = client.post("/api/v1/resumes/upload", headers=headers, files=files)

    assert resp.status_code == 201, f"Resume upload failed: {resp.text}"
    resume = resp.json()
    resume_id = resume["id"]
    print(f"Resume uploaded successfully! ID: {resume_id}, Status: {resume['status']}")

    # Print detected skills if parsed
    if resume.get("parsed_content") and resume["parsed_content"].get("skills"):
        skills = resume["parsed_content"]["skills"]
        print(f"  Detected skills: {', '.join(skills[:8])}{'...' if len(skills) > 8 else ''}")

    print("\n6. Testing Duplicate Resume Detection...")
    with open(sample_docx_path, "rb") as f:
        files = {"file": ("john_doe_resume.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        dup_resp = client.post("/api/v1/resumes/upload", headers=headers, files=files)
    
    assert dup_resp.status_code == 400
    assert "Duplicate resume detected" in dup_resp.json()["detail"]
    print("Duplicate detection verified! Blocked upload of identical file.")



    print("\n7. Testing Resume History Listing...")
    resp = client.get("/api/v1/resumes/", headers=headers)
    assert resp.status_code == 200
    resumes_list = resp.json()
    print(f"Resume History count: {len(resumes_list)}")

    print("\n8. Testing Resume Parse API...")
    resp = client.post(f"/api/v1/resumes/{resume_id}/parse", headers=headers)
    assert resp.status_code == 200
    print("Parsed output status:", resp.json()["status"])

    print("\n9. Testing Resume Download...")
    resp = client.get(f"/api/v1/resumes/{resume_id}/download", headers=headers)
    assert resp.status_code == 200
    print("Download response status:", resp.status_code, len(resp.content), "bytes")

    print("\n10. Testing Resume Deletion...")
    resp = client.delete(f"/api/v1/resumes/{resume_id}", headers=headers)
    assert resp.status_code == 204
    print("Resume deleted successfully!")

    print("\nALL BACKEND TESTS PASSED FOR WEEKS 1 TO 5!")

if __name__ == "__main__":
    test_full_flow()
