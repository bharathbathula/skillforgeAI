import os
import sys
import warnings
import tempfile
from docx import Document as DocxDocument

# Suppress starlette/httpx deprecation warning
warnings.filterwarnings("ignore", category=DeprecationWarning)

from fastapi.testclient import TestClient

# Add current directory to path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from app.main import app

client = TestClient(app)

def create_sample_resume_docx():
    temp_dir = tempfile.gettempdir()
    file_path = os.path.join(temp_dir, "sample_ats_resume_test.docx")

    doc = DocxDocument()
    doc.add_heading("Alex Candidate", 0)
    doc.add_paragraph("Email: alex.ats@example.com | Phone: (555) 123-4567 | Location: New York, NY")
    doc.add_paragraph("LinkedIn: linkedin.com/in/alexcandidate | GitHub: github.com/alexcandidate")
    
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph("High-performing Software Engineer with expertise in building scalable REST APIs using Python, FastAPI, React, and PostgreSQL. Experienced in database optimization and cloud containerization.")
    
    doc.add_heading("Skills", level=1)
    doc.add_paragraph("Python, JavaScript, TypeScript, FastAPI, React, PostgreSQL, SQLite, Git, Docker, REST API, HTML, CSS, SQL, PyTest")
    
    doc.add_heading("Work Experience", level=1)
    doc.add_paragraph("Software Engineer Intern - Tech Solutions Corp (2023 - 2024)")
    doc.add_paragraph("- Engineered high-speed RESTful APIs using Python and FastAPI.")
    doc.add_paragraph("- Designed PostgreSQL database schemas and optimized complex SQL queries.")

    doc.add_heading("Projects", level=1)
    doc.add_paragraph("SkillForge AI - Resume Analyzer & ATS Scoring System")
    doc.add_paragraph("- Built full-stack platform using Python, FastAPI, PostgreSQL, and React.")
    doc.add_paragraph("- Implemented text extraction, skill matching, and vector embedding similarity algorithms.")

    doc.add_heading("Education", level=1)
    doc.add_paragraph("B.Tech in Computer Science - State University (2020 - 2024)")

    doc.save(file_path)
    return file_path

def test_full_pipeline():
    print("\n--- STARTING SKILLFORGE AI FULL PIPELINE TEST ---")

    # 1. Health Check
    res = client.get("/")
    assert res.status_code == 200
    print("1. Health Endpoint [OK] PASS:", res.json()["status"])

    # 2. Register & Login
    email = "alex.ats@example.com"
    pwd = "SecurePassword123!"
    client.post("/api/v1/auth/register", json={"email": email, "password": pwd, "full_name": "Alex Candidate"})
    
    login_res = client.post("/api/v1/auth/login/json", json={"email": email, "password": pwd})
    assert login_res.status_code == 200
    token = login_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    print("2. Authentication (Register + Login) [OK] PASS")

    # 3. Resume Upload
    docx_path = create_sample_resume_docx()
    with open(docx_path, "rb") as f:
        up_res = client.post("/api/v1/resumes/upload", files={"file": ("sample_ats_resume_test.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}, headers=headers)
    
    assert up_res.status_code == 201
    resume_id = up_res.json()["id"]
    print(f"3. Resume DOCX Upload [OK] PASS: Resume ID {resume_id}")

    # 4. Job Description Input
    jd_payload = {
        "title": "Backend Software Engineer",
        "description": "We are seeking a Backend Software Engineer skilled in Python, FastAPI, PostgreSQL, REST APIs, and Docker. Experience building database schemas, writing SQL queries, and working with React is preferred.",
        "resume_id": resume_id
    }
    jd_res = client.post("/api/v1/jobdescription/", json=jd_payload, headers=headers)
    assert jd_res.status_code == 201
    job_id = jd_res.json()["id"]
    print(f"4. Job Description Post [OK] PASS: Job ID {job_id}")

    # 5. Execute Full Analysis & ATS Pipeline
    analysis_payload = {
        "resume_id": resume_id,
        "job_description_id": job_id
    }
    analysis_res = client.post(f"/api/v1/analysis/{resume_id}", json=analysis_payload, headers=headers)
    if analysis_res.status_code != 201:
        print("Analysis Error response:", analysis_res.status_code, analysis_res.json())
    assert analysis_res.status_code == 201
    report = analysis_res.json()["report_details"]
    overall_score = analysis_res.json()["overall_score"]
    
    print(f"5. Complete Analysis & Matching Pipeline [OK] PASS")
    print(f"   -> Overall ATS Score: {overall_score} / 100")
    print(f"   -> Skill Score: {report['score_breakdown']['skill_score']}%")
    print(f"   -> Keyword Score: {report['score_breakdown']['keyword_score']}%")
    print(f"   -> Project Score: {report['score_breakdown']['project_score']}%")
    print(f"   -> Matched Skills: {', '.join(report['skill_analysis']['matched_skills'])}")

    # 6. Retrieve ATS Full Report API
    ats_res = client.get(f"/api/v1/ats/{resume_id}/report", headers=headers)
    assert ats_res.status_code == 200
    print("6. GET /ats/{resume_id}/report API [OK] PASS")

    print("\nALL PIPELINE TESTS PASSED CLEANLY!\n")

if __name__ == "__main__":
    test_full_pipeline()
